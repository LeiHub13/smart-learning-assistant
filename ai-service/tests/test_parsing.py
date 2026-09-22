"""文档解析与分块测试（对应自 Java 侧迁移过来的 DocumentParser 行为）。"""
import io

import pytest

from app import config, parsing


def test_parse_plain_text():
    data = "Java 是一门面向对象的编程语言。".encode("utf-8")
    assert parsing.parse("note.txt", "text/plain", data) == "Java 是一门面向对象的编程语言。"


def test_parse_markdown_by_extension():
    data = "# 标题\n正文".encode("utf-8")
    assert parsing.parse("doc.md", None, data) == "# 标题\n正文"


def test_reject_non_utf8_text():
    data = "中文".encode("gbk")
    with pytest.raises(parsing.ParseError, match="UTF-8"):
        parsing.parse("note.txt", "text/plain", data)


def test_reject_binary_masquerading_as_txt():
    data = b"\xd0\xcf\x11\xe0\x00\x00\x00\x00"  # OLE 复合文档头，含 NUL
    with pytest.raises(parsing.ParseError, match="不是可解析的文本"):
        parsing.parse("evil.txt", "text/plain", data)


def test_reject_legacy_doc_with_actionable_message():
    with pytest.raises(parsing.ParseError, match="另存为 .docx"):
        parsing.parse("旧讲义.doc", "application/msword", b"\xd0\xcf\x11\xe0")


def test_reject_empty_extraction():
    with pytest.raises(parsing.ParseError, match="未能从该文档提取到文字"):
        parsing.parse("blank.md", "text/markdown", b"   \n  ")


def test_parse_docx(tmp_path):
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("第一章 绪论")
    doc.add_paragraph("向量库使用 Chroma。")
    doc.save(buf)
    text = parsing.parse("讲义.docx",
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         buf.getvalue())
    assert "第一章 绪论" in text
    assert "Chroma" in text


def test_parse_pdf_blank_page_reports_readable_error():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(parsing.ParseError, match="未能从该文档提取到文字"):
        # 空白 PDF 提不出文字，应给可读原因（提示需 OCR）而非静默产出 0 chunk
        parsing.parse("x.pdf", "application/pdf", buf.getvalue())


def test_to_chunks_matches_legacy_java_split():
    text = "a" * 450
    chunks = parsing.to_chunks(text)
    assert [len(c) for c in chunks] == [200, 200, 50]
    assert "".join(chunks) == text


def test_to_chunks_respects_config_size(monkeypatch):
    monkeypatch.setattr(config, "RAG_CHUNK_SIZE", 10)
    assert parsing.to_chunks("0123456789abc") == ["0123456789", "abc"]


def test_extension_parsing():
    assert parsing.extension("A.B.PDF") == "pdf"
    assert parsing.extension("noext") == ""
    assert parsing.extension(None) == ""
